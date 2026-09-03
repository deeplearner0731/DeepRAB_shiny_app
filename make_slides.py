"""
Build the 2-slide deck 'DeepRAB_vs_SeqBATTing.pptx'.

Standalone utility, not part of the app.  Needs python-pptx, which is
deliberately NOT installed in .venv (that venv is what gets deployed):

    python -m pip install --target /tmp/pptxlib python-pptx
    PYTHONPATH=/tmp/pptxlib python make_slides.py
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = "DeepRAB_vs_SeqBATTing.pptx"

NAVY = RGBColor(0x1F, 0x35, 0x64)
INK = RGBColor(0x22, 0x26, 0x2B)
GREY = RGBColor(0x5C, 0x63, 0x6B)
RULE = RGBColor(0xD6, 0xDB, 0xE1)
BAND = RGBColor(0xEE, 0xF1, 0xF5)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

MARGIN = Inches(0.55)
WIDTH = prs.slide_width - 2 * MARGIN


def add_slide():
    return prs.slides.add_slide(BLANK)


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size, *, bold=False, color=INK, first=False,
         space_before=0, space_after=2, indent=0, italic=False, bullet=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.level = indent
    body = f"{bullet} {text}" if bullet else text
    run = p.add_run()
    run.text = body
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return p


def header(slide, title, subtitle=None):
    tf = textbox(slide, MARGIN, Inches(0.34), WIDTH, Inches(0.8))
    para(tf, title, 26, bold=True, color=NAVY, first=True, space_after=2)
    if subtitle:
        para(tf, subtitle, 12.5, color=GREY, italic=True)
    line = slide.shapes.add_shape(1, MARGIN, Inches(1.12), WIDTH, Emu(9525))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    line.shadow.inherit = False


def footnote(slide, text, top=Inches(6.95)):
    tf = textbox(slide, MARGIN, top, WIDTH, Inches(0.4))
    para(tf, text, 9.5, color=GREY, italic=True, first=True)


# ============================================================ slide 1 #
s1 = add_slide()
header(
    s1,
    "DeepRAB: what questions does it address?",
    "Exploratory analysis of a randomized trial — baseline covariates only; "
    "continuous, binary, or time-to-event endpoints",
)

COL_W = Inches(3.95)
GAP = Inches(0.24)
TOP = Inches(1.45)
COL_H = Inches(3.55)

questions = [
    (
        "Q1  Which variables are PREDICTIVE (effect modifiers), not merely prognostic?",
        [
            "Output: ranked marker list with a selection probability per variable",
            "The A-/R-learning loss targets the treatment × covariate contrast "
            "directly; baseline effect is removed by a cross-fitted nuisance model",
            "So a strongly prognostic-only marker does not rise in the ranking",
        ],
    ),
    (
        "Q2  Which patients benefit, and by how much?",
        [
            "Output: per-patient contrast τ̂(x) on the native scale — "
            "mean difference, log OR, or log HR",
            "Predicted-benefit subgroup = { τ̂(x) > 0 }",
            "Subgroup size and treatment effect with bootstrap CI reported for "
            "benefit vs. no-benefit groups",
            "No functional form pre-specified: nonlinear effects and "
            "marker × marker interactions are learned",
        ],
    ),
    (
        "Q3  Is the signal real, or an artifact of searching?",
        [
            "One fit / validation / test split — the reported subgroup effect "
            "comes from held-out patients",
            "Nuisance models cross-fitted: no patient contributes to its own "
            "adjustment",
            "Hyperparameters picked by MEAN validation loss across restarts, not "
            "the single best fit (avoids winner's-curse bias)",
            "Final score is an ensemble of the top configurations",
        ],
    ),
]

for i, (q_title, bullets) in enumerate(questions):
    left = MARGIN + i * (COL_W + GAP)
    card = s1.shapes.add_shape(1, left, TOP, COL_W, COL_H)
    card.fill.solid()
    card.fill.fore_color.rgb = BAND
    card.line.color.rgb = RULE
    card.shadow.inherit = False

    tf = textbox(s1, left + Inches(0.2), TOP + Inches(0.18),
                 COL_W - Inches(0.4), COL_H - Inches(0.36))
    para(tf, q_title, 14, bold=True, color=NAVY, first=True, space_after=8)
    for b in bullets:
        para(tf, b, 11.5, bullet="•", space_after=6)

BOTTOM = TOP + COL_H + Inches(0.32)
box = s1.shapes.add_shape(1, MARGIN, BOTTOM, WIDTH, Inches(1.05))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0xFF, 0xF7, 0xE6)
box.line.color.rgb = RGBColor(0xE2, 0xC4, 0x7A)
box.shadow.inherit = False

tf = textbox(s1, MARGIN + Inches(0.2), BOTTOM + Inches(0.16),
             WIDTH - Inches(0.4), Inches(0.8))
para(tf, "What DeepRAB does NOT deliver", 13, bold=True, color=NAVY,
     first=True, space_after=4)
para(tf,
     "A signature with biomarker cutoffs — and not a confirmatory result. "
     "Findings are hypothesis-generating and require independent validation "
     "(see slide 2 for how Sequential BATTing supplies the cutoff).",
     12)

# ============================================================ slide 2 #
s2 = add_slide()
header(
    s2,
    "Key difference from Sequential BATTing",
    "Sequential BATTing: Huang, Sun, Trow, Chatterjee, Chakravartty, Tian, "
    "Devanarayan, Statistics in Medicine 2017 (R package SubgrpID)",
)

rows = [
    ("Primary output",
     "Signature rule with explicit cutoffs, e.g. (BM1 ≥ c₁) & "
     "(BM2 ≤ c₂) → signature +/−",
     "Continuous per-patient contrast τ̂(x) + ranked marker list; "
     "subgroup = τ̂(x) > 0"),
    ("How the subgroup\nis built",
     "Forward stepwise: pick the marker/threshold maximizing a treatment-effect "
     "(interaction) statistic; threshold stabilized by bootstrap aggregation of "
     "tree splits",
     "Neural net with a concrete (differentiable) feature-selection layer trained "
     "on an A-/R-learning loss; subgroup falls out of the sign of the fitted "
     "contrast"),
    ("Subgroup shape",
     "Axis-aligned box — AND of monotone thresholds",
     "Smooth, possibly nonlinear boundary in the selected markers"),
    ("Prognostic effect",
     "Not explicitly orthogonalized",
     "Removed by construction (residualized / offset nuisance)"),
    ("Scaling in # of\ncandidate markers",
     "Greedy stepwise; practical for a modest panel",
     "Built for wider panels; selection layer picks K markers jointly"),
    ("Deployability",
     "High — cutoffs go straight into an enrichment criterion or assay spec",
     "Lower — scoring a new patient requires the fitted model"),
    ("Endpoints",
     "Continuous / binary / survival",
     "Continuous / binary / Cox"),
    ("Main risk",
     "Greedy search + rectangular shape can miss diffuse or interacting signals; "
     "cutoff optimism",
     "Black-box score; more data-hungry; no ready-made cutoff"),
]

T_TOP = Inches(1.4)
T_LEFT = MARGIN
T_W = Inches(7.55)
label_w, col_w = Inches(1.55), Inches(3.0)

table = s2.shapes.add_table(len(rows) + 1, 3, T_LEFT, T_TOP, T_W,
                            Inches(0.3) * (len(rows) + 1)).table
table.columns[0].width = label_w
table.columns[1].width = col_w
table.columns[2].width = col_w
table.first_row = True

heads = ("", "Sequential BATTing", "DeepRAB")
for j, h in enumerate(heads):
    cell = table.cell(0, j)
    cell.text = h
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY

for i, (label, batt, deep) in enumerate(rows, start=1):
    table.rows[i].height = Inches(0.44)
    for j, text in enumerate((label, batt, deep)):
        cell = table.cell(i, j)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = (BAND if i % 2 else
                                    RGBColor(0xFF, 0xFF, 0xFF))
        for p in cell.text_frame.paragraphs:
            p.space_after = Pt(0)
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.name = "Calibri"
                run.font.bold = (j == 0)
                run.font.color.rgb = NAVY if j == 0 else INK

R_LEFT = T_LEFT + T_W + Inches(0.3)
R_W = prs.slide_width - MARGIN - R_LEFT

panel = s2.shapes.add_shape(1, R_LEFT, T_TOP, R_W, Inches(4.55))
panel.fill.solid()
panel.fill.fore_color.rgb = BAND
panel.line.color.rgb = RULE
panel.shadow.inherit = False

tf = textbox(s2, R_LEFT + Inches(0.22), T_TOP + Inches(0.18),
             R_W - Inches(0.44), Inches(4.2))
para(tf, "One-line framing", 13, bold=True, color=NAVY, first=True,
     space_after=4)
para(tf,
     "Sequential BATTing answers “what is the rule?”  DeepRAB answers "
     "“for whom, and how strongly, does the effect vary?”",
     11.5, space_after=12)

para(tf, "How we would use the two together", 13, bold=True, color=NAVY,
     space_after=5)
joint = [
    ("Mutual validation (primary use).", " Run both; concordance in selected "
     "markers, direction of benefit and subgroup membership (overlap %, "
     "κ) strengthens an exploratory finding — divergence flags a "
     "fragile one."),
    ("Two-stage pipeline.", " DeepRAB screens and ranks the panel (nonlinear, "
     "prognostic-adjusted); Sequential BATTing then derives the deployable "
     "cutoff on the top 1–3 markers."),
    ("Common report card.", " On the same held-out patients, compare subgroup "
     "size, effect estimate + CI, and prevalence for each method's subgroup."),
    ("Exploratory by design.", " Reporting pre-specified before the subgroup "
     "effects are looked at."),
]
for k, (lead, rest) in enumerate(joint, start=1):
    p = tf.add_paragraph()
    p.space_after = Pt(7)
    r1 = p.add_run()
    r1.text = f"{k}.  {lead}"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = INK
    r1.font.name = "Calibri"
    r2 = p.add_run()
    r2.text = rest
    r2.font.size = Pt(11)
    r2.font.color.rgb = INK
    r2.font.name = "Calibri"

CAVEAT_TOP = T_TOP + Inches(4.55) + Inches(0.2)
box = s2.shapes.add_shape(1, MARGIN, CAVEAT_TOP, WIDTH, Inches(0.95))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0xFF, 0xF7, 0xE6)
box.line.color.rgb = RGBColor(0xE2, 0xC4, 0x7A)
box.shadow.inherit = False

tf = textbox(s2, MARGIN + Inches(0.18), CAVEAT_TOP + Inches(0.13),
             WIDTH - Inches(0.36), Inches(0.7))
p = para(tf, "Caveat on “DeepRAB gives no cutoff”:  ", 11, bold=True,
         color=NAVY, first=True, space_after=0)
r = p.add_run()
r.text = ("τ̂(x) = 0 is itself a boundary, so for a single dominant "
          "marker an implied cutoff can be read off where τ̂ crosses zero "
          "along that marker (others at their medians), or τ̂ can be "
          "distilled with a shallow tree. But that boundary generally depends on "
          "the other markers and is post-hoc — it is not an assay-grade "
          "cutoff the way a BATTing threshold is, which is exactly why point 2 "
          "hands the cutoff to Sequential BATTing.")
r.font.size = Pt(11)
r.font.color.rgb = INK
r.font.name = "Calibri"

prs.save(OUT)
print(f"wrote {OUT}")
